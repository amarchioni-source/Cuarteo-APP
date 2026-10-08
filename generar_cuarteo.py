"""
Generador del CUARTEO final (Anexo Cuarteo y Troceo, circular 3510) a partir
de la bajada del sistema ("cuarteo de sistema").

Lógica
------
* Formulario (filas 17 en adelante): tropas, kg de ingreso y medias salen de
  la bajada. Cambian solo la fecha de faena y el número de serie.
* Bloque de cálculo (debajo del formulario, fuera del área de impresión), por
  cada tropa y para 5 productos:
      cantidad = medias de la tropa
      kg       = cantidad x peso promedio por media
* El peso promedio por media se calcula SIEMPRE desde la bajada del día:
      kg/media = kg de la categoría en la bajada (todas las tropas)
                 / total de medias del día
  El cálculo queda a la vista en la hoja "Peso promedio" (producto de la
  bajada -> categoría final, editable) y la columna J del detalle lo toma de
  ahí por fórmula.

Uso por línea de comandos
-------------------------
    python generar_cuarteo.py BAJADA.xlsx PLANTILLA.xlsx SALIDA.xlsx \
           [--fecha-faena 2026-10-05] [--serie 2291]
"""
import argparse
import copy
import datetime as dt
import io
import re
import sys
import zipfile

import openpyxl
import PIL  # noqa: F401  (sin Pillow, openpyxl descarta el logo de SENASA sin avisar)
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.properties import CalcProperties

HOJA_PESO = "Peso promedio"
PRODUCTOS = ["Asado CV", "Delantero SF", "Entraña Fina", "Falda", "Pistola "]
SIN_CAT = "No va"

# Reglas producto de la bajada -> categoría final (gana la primera que coincide).
# Se pueden corregir a mano en la hoja "Peso promedio" sin tocar el código.
REGLAS = [
    # Asado: con vacío (CV) y sin vacío (SV, banderita) + el vacío colgado que sale aparte
    (r"^ASADO", "Asado CV"),
    (r"^VACIO COLGADO", "Asado CV"),
    # Delantero sin falda: delantero (completo, musulmán, SF) + pecho CB SF
    (r"^DELANTERO", "Delantero SF"),
    (r"^PECHO CB SF", "Delantero SF"),
    (r"^ENTRAÑA", "Entraña Fina"),
    (r"^FALDA", "Falda"),
    # Pistola: pistola entera + pistola trozada en R&L 2 costillas + Rueda
    (r"^(PISTOLA|R&L|RUEDA)", "Pistola "),
]
# Productos que se sabe que no van al final. Cualquier otro producto sin regla
# genera un aviso ("producto nuevo") para que se revise.
NO_VA_CONOCIDO = r"^(DECOMISO|RECORTE|FALSA ENTRA|VARIOS)"

FIRST = 18                                  # primera fila de tropas
COLS_MERGE = [("C", "F"), ("G", "H"), ("I", "K"), ("M", "Q"), ("S", "W")]
MAX_TROPAS = 200


class BajadaInvalida(ValueError):
    """La bajada no tiene el formato esperado (mensaje apto para mostrar)."""


def categoria(nombre):
    for rx, cat in REGLAS:
        if re.search(rx, nombre.upper()):
            return cat
    return SIN_CAT


def dia_habil_anterior(d):
    d -= dt.timedelta(days=1)
    while d.weekday() >= 5:
        d -= dt.timedelta(days=1)
    return d


# ------------------------------------------------------------------ lectura
def leer_bajada(src):
    """Lee la bajada. `src` puede ser una ruta o un archivo en memoria."""
    try:
        ws = openpyxl.load_workbook(src, data_only=True).active
    except Exception as e:                                   # noqa: BLE001
        raise BajadaInvalida(
            "No se pudo abrir el archivo. Verificá que sea un .xlsx válido.") from e

    if ("CUARTEO" not in str(ws["E4"].value or "").upper()
            or "tropa" not in str(ws["C17"].value or "").lower()):
        raise BajadaInvalida(
            "El archivo no tiene el formato de la bajada de cuarteo "
            "(se esperaba el Anexo Cuarteo y Troceo, hoja PD_AnexoIII).")

    m = re.search(r"D[ií]a\s+(\d+)\s+Mes\s+(\d+)\s+A[ñn]o\s+(\d+)",
                  str(ws["O14"].value or ""))
    if not m:
        raise BajadaInvalida("No se encontró la fecha del programa (celda O14).")
    fecha = dt.date(int(m[3]), int(m[2]), int(m[1]))

    tropas, r = [], FIRST
    while True:
        if r > FIRST + MAX_TROPAS:
            raise BajadaInvalida("No se encontró la fila TOTAL de la tabla de tropas.")
        if str(ws.cell(r, 9).value).strip() == "TOTAL":
            break
        c = ws.cell(r, 3).value
        if c in (None, ""):
            raise BajadaInvalida(f"Fila {r}: falta el número de tropa antes del TOTAL.")
        kg, medias = ws.cell(r, 12).value, ws.cell(r, 13).value
        if not isinstance(kg, (int, float)) or not isinstance(medias, (int, float)):
            raise BajadaInvalida(f"Tropa {c} (fila {r}): kg de ingreso o medias no numéricos.")
        tropas.append(dict(tropa=str(c).strip(), materia=ws.cell(r, 9).value,
                           kg=kg, medias=medias, obs=ws.cell(r, 19).value or "-"))
        r += 1
    if not tropas:
        raise BajadaInvalida("La bajada no tiene tropas.")
    total_medias_bajada = ws.cell(r, 13).value

    por_tropa, tot_sis, cur = {}, {}, None
    for rr in range(r + 1, ws.max_row + 1):
        b = ws.cell(rr, 2).value
        if b not in (None, ""):
            cur = str(b).strip()
        h, k, n = ws.cell(rr, 8).value, ws.cell(rr, 11).value, ws.cell(rr, 14).value
        if str(k).strip() == "TOTAL :" and cur:
            tot_sis[cur] = n
        elif h not in (None, "") and cur and isinstance(k, (int, float)):
            nombre = re.sub(r"\s+", " ", str(h)).strip()
            p = por_tropa.setdefault(cur, {}).setdefault(nombre, [0, 0.0])
            p[0] += k
            p[1] += n or 0
    if not por_tropa:
        raise BajadaInvalida("La bajada no tiene el detalle de productos por tropa.")

    agg = {}
    for prods in por_tropa.values():
        for nombre, (pz, kg) in prods.items():
            a = agg.setdefault(nombre, [0, 0.0])
            a[0] += pz
            a[1] += kg
    detalle = sorted((n, v[0], round(v[1], 4)) for n, v in agg.items())

    return dict(tropas=tropas, total_medias_bajada=total_medias_bajada,
                fecha=fecha, tot_sis=tot_sis, detalle=detalle,
                tropas_detalle=set(por_tropa))


# ----------------------------------------------------------------- análisis
def analizar(d):
    """Cálculo en Python (el libro lo repite con fórmulas). Devuelve resumen."""
    tropas, detalle = d["tropas"], d["detalle"]
    medias = sum(t["medias"] for t in tropas)
    avisos = []

    if d["total_medias_bajada"] not in (None, "") and d["total_medias_bajada"] != medias:
        avisos.append(f"El total de medias de la bajada ({d['total_medias_bajada']}) "
                      f"no coincide con la suma de las tropas ({medias}).")
    ids = {t["tropa"] for t in tropas}
    sin_det = sorted(ids - d["tropas_detalle"])
    if sin_det:
        avisos.append("Tropas sin detalle de productos en la bajada: " + ", ".join(sin_det) + ".")
    sin_cab = sorted(d["tropas_detalle"] - ids)
    if sin_cab:
        avisos.append("Tropas con detalle pero ausentes en la tabla de tropas: "
                      + ", ".join(sin_cab) + ".")

    kg_cat = {c: 0.0 for c in PRODUCTOS}
    no_va = 0.0
    nuevos = []
    for nombre, _pz, kg in detalle:
        c = categoria(nombre)
        if c == SIN_CAT:
            no_va += kg
            if not re.search(NO_VA_CONOCIDO, nombre.upper()):
                nuevos.append(nombre)
        else:
            kg_cat[c] += kg
    if nuevos:
        avisos.append("Productos nuevos sin categoría (quedaron fuera del final; "
                      "asignarlos en la hoja «Peso promedio» si corresponde): "
                      + "; ".join(nuevos) + ".")

    cats = [dict(nombre=c.strip(), kg=kg_cat[c], kg_media=kg_cat[c] / medias)
            for c in PRODUCTOS]
    suma_j = sum(c["kg_media"] for c in cats)
    filas = []
    for t in tropas:
        kf = t["medias"] * suma_j
        ks = d["tot_sis"].get(t["tropa"])
        filas.append(dict(tropa=t["tropa"], medias=t["medias"], kg_ingreso=t["kg"],
                          kg_final=kf, kg_sistema=ks,
                          dif_pct=(kf / ks - 1) * 100 if ks else None))
    return dict(n_tropas=len(tropas), medias=medias,
                kg_ingreso=sum(t["kg"] for t in tropas),
                kg_sistema=sum(v for v in d["tot_sis"].values() if v),
                kg_final=sum(c["kg"] for c in cats), no_va_kg=no_va,
                categorias=cats, tropas=filas, avisos=avisos)


# ---------------------------------------------------------- formato (helpers)
def _snap(ws, row, c1=2, c2=28):
    return {c: (ws.cell(row, c).value, copy.copy(ws.cell(row, c)._style))
            for c in range(c1, c2 + 1)}


def _pegar(ws, row, snap, valores):
    for c, (_v, st) in snap.items():
        cell = ws.cell(row, c)
        cell._style = copy.copy(st)
        cell.value = valores.get(c)


def _fila_total_plantilla(ws):
    for r in range(FIRST, FIRST + MAX_TROPAS):
        if str(ws.cell(r, 9).value).strip() == "TOTAL":
            return r
    raise ValueError("La plantilla no tiene fila TOTAL.")


def _hoja_peso(wb, detalle, row_total, ws_form):
    wp = wb.create_sheet(HOJA_PESO)
    f = lambda **k: Font(name="Segoe UI", size=9, **k)            # noqa: E731
    thin = Side(style="thin", color="BFBFBF")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    hdr = PatternFill("solid", fgColor="D9E1F2")
    wp["A1"] = "Peso promedio por media — calculado desde la bajada del sistema"
    wp["A1"].font = f(bold=True)
    for col, t in zip("ABCDFGHI", ["CATEGORÍA FINAL", "KG BAJADA", "MEDIAS DEL DÍA",
                                   "KG / MEDIA", "PRODUCTO (BAJADA)", "PIEZAS", "KG",
                                   "CATEGORÍA FINAL (editable)"]):
        c = wp[f"{col}3"]
        c.value = t
        c.font, c.fill, c.border = f(bold=True), hdr, box
    fin = 3 + len(detalle)
    for i, (n, pz, kg) in enumerate(detalle):
        r = 4 + i
        wp[f"F{r}"], wp[f"G{r}"], wp[f"H{r}"], wp[f"I{r}"] = n, pz, kg, categoria(n)
        for col in "FGHI":
            wp[f"{col}{r}"].font = f(color="0000FF" if col in "GHI" else "000000")
            wp[f"{col}{r}"].border = box
        wp[f"H{r}"].number_format = "#,##0.00"
    for j, cat in enumerate(PRODUCTOS):
        r = 4 + j
        wp[f"A{r}"] = cat
        wp[f"B{r}"] = f"=SUMIF($I$4:$I${fin},A{r},$H$4:$H${fin})"
        wp[f"C{r}"] = f"='{ws_form}'!$M${row_total}"
        wp[f"D{r}"] = f"=IF(C{r}=0,0,B{r}/C{r})"
        wp[f"B{r}"].number_format = "#,##0.00"
        wp[f"D{r}"].number_format = "0.0000"
        for col in "ABCD":
            wp[f"{col}{r}"].font, wp[f"{col}{r}"].border = f(), box
    wp["A10"] = "Kg de la bajada que no entran al final"
    wp["B10"] = f'=SUMIF($I$4:$I${fin},"{SIN_CAT}",$H$4:$H${fin})'
    wp["B10"].number_format = "#,##0.00"
    wp["A12"] = ("Categoría 'No va' = producto de la bajada que no se informa en el final. "
                 "Para cambiar la asignación, editar la columna I; el KG/MEDIA se recalcula solo.")
    crit = ["Criterios de asignación:",
            "Asado CV = ASADO CV + ASADO SV (banderita) + VACÍO COLGADO (sale aparte del asado SV)",
            "Delantero SF = DELANTERO (completo, musulmán, SF) + PECHO CB SF (musulmán, rechazo)",
            "Entraña Fina = ENTRAÑA (AA, World Class). Falsa entraña no va",
            "Falda = FALDA",
            "Pistola = PISTOLA 2C + R&L 2 COST + RUEDA (la pistola trozada)",
            "No va = Decomiso, Recorte, Falsa entraña y Varios (no son producto de cuarteo)"]
    for i, t in enumerate(crit):
        wp[f"A{14 + i}"] = t
        wp[f"A{14 + i}"].font = f(bold=(i == 0))
    for a in ("A10", "B10"):
        wp[a].font = f()
    wp["A12"].font = f(italic=True)
    for col, w in zip("ABCDEFGHI", [34, 14, 16, 12, 3, 44, 9, 12, 26]):
        wp.column_dimensions[col].width = w
    wp.sheet_view.showGridLines = False


def _poner_fecha(cell, fecha):
    """Cambia día, mes y año conservando los tramos de formato del texto."""
    v = cell.value
    nums = iter([str(fecha.day), str(fecha.month), str(fecha.year)])
    if isinstance(v, CellRichText):
        nuevo = []
        for el in v:                      # lista nueva: asignar por índice fusiona tramos
            if isinstance(el, TextBlock) and el.text.strip().isdigit():
                n = next(nums, None)
                el = TextBlock(el.font, n) if n is not None else el
            nuevo.append(el)
        cell.value = CellRichText(nuevo)
    else:
        d, m, y = nums
        cell.value = f"Fecha: Día {d} Mes {m} Año {y}"


def _poner_serie(cell, serie):
    """Cambia el número final de «SERIE C N° nnnn» conservando el formato."""
    v = cell.value
    if isinstance(v, CellRichText):
        nuevo = list(v)
        el = nuevo[-1]
        if isinstance(el, TextBlock):
            nuevo[-1] = TextBlock(el.font, re.sub(r"\d+\s*$", str(serie), el.text))
        else:
            nuevo[-1] = re.sub(r"\d+\s*$", str(serie), el)
        cell.value = CellRichText(nuevo)
    else:
        cell.value = re.sub(r"\d+\s*$", str(serie), str(v))


def _texto(cell):
    v = cell.value
    return "".join(el.text if isinstance(el, TextBlock) else el for el in v) \
        if isinstance(v, CellRichText) else str(v)


def _filas_autoaltura(plantilla):
    """Filas de la plantilla con altura guardada pero sin 'altura personalizada'
    (Excel las ajusta al contenido). openpyxl les agrega customHeight al guardar."""
    if hasattr(plantilla, "seek"):
        plantilla.seek(0)
    with zipfile.ZipFile(plantilla) as z:
        x = z.read("xl/worksheets/sheet1.xml").decode("utf8")
    filas = []
    for m in re.finditer(r"<row ([^>]*)>", x):
        a = dict(re.findall(r'(\w+)="([^"]*)"', m.group(1)))
        if "ht" in a and a.get("customHeight") != "1" and int(a["r"]) < FIRST:
            filas.append(int(a["r"]))
    if hasattr(plantilla, "seek"):
        plantilla.seek(0)
    return filas


def _quitar_custom_height(datos, filas):
    if not filas:
        return datos
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(datos)) as zin, \
            zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            blob = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                x = blob.decode("utf8")
                for r in filas:
                    x = re.sub(r'(<row r="%d"[^>]*?) customHeight="1"' % r, r"\1", x, count=1)
                blob = x.encode("utf8")
            zout.writestr(item, blob)
    return out.getvalue()


# ---------------------------------------------------------------- generación
def generar(bajada, plantilla, salida, fecha_faena=None, serie=None):
    """
    bajada, plantilla: ruta o archivo en memoria.  salida: ruta o archivo en memoria.
    Devuelve el resumen (ver `analizar`) más fecha, fecha_faena y serie.
    """
    d = leer_bajada(bajada)
    res = analizar(d)
    tropas, detalle, fecha = d["tropas"], d["detalle"], d["fecha"]

    filas_auto = _filas_autoaltura(plantilla)
    wb = openpyxl.load_workbook(plantilla, rich_text=True)
    ws = wb.active
    t0 = _fila_total_plantilla(ws)                  # fila TOTAL en la plantilla

    # estilos de la plantilla (todo relativo a su fila TOTAL)
    s_dato, s_total = _snap(ws, FIRST), _snap(ws, t0)
    cola = {o: _snap(ws, t0 + o) for o in range(1, 6)}
    alt_cola = {o: ws.row_dimensions[t0 + o].height for o in range(0, 6)}
    alt_dato = ws.row_dimensions[FIRST].height
    cola_merges = [m for m in ws.merged_cells.ranges if t0 < m.min_row <= t0 + 5]
    h0 = t0 + 6                                     # encabezado del bloque de cálculo
    s_h, s_p1 = _snap(ws, h0, 6, 10), _snap(ws, h0 + 1, 6, 10)
    s_p, s_sum = _snap(ws, h0 + 2, 6, 10), _snap(ws, h0 + 7, 6, 10)
    alt_det = ws.row_dimensions[h0 + 1].height

    # limpiar todo desde la primera fila de tropas
    for m in [m for m in ws.merged_cells.ranges if m.min_row >= FIRST]:
        ws.unmerge_cells(str(m))
    for k in [k for k in ws._cells if k[0] >= FIRST]:
        del ws._cells[k]
    for k in [k for k in ws.row_dimensions if k >= FIRST]:
        del ws.row_dimensions[k]

    ff = fecha_faena or dia_habil_anterior(fecha)
    _poner_fecha(ws["O14"], fecha)
    if serie is not None:
        _poner_serie(ws["P6"], serie)
    else:
        res["avisos"].append("No se indicó el número de serie: quedó el de la plantilla. Revisarlo.")
    m_serie = re.search(r"(\d+)\s*$", _texto(ws["P6"]))

    # formulario: una fila por tropa
    r = FIRST
    for t in tropas:
        _pegar(ws, r, s_dato, {3: t["tropa"], 7: dt.datetime.combine(ff, dt.time()),
                               9: t["materia"], 12: t["kg"], 13: t["medias"], 19: t["obs"]})
        for a, b in COLS_MERGE:
            ws.merge_cells(f"{a}{r}:{b}{r}")
        if alt_dato:
            ws.row_dimensions[r].height = alt_dato
        r += 1
    row_total = r
    _pegar(ws, row_total, s_total, {9: "TOTAL", 13: res["medias"]})
    shift = row_total - t0
    for o in range(0, 6):
        if alt_cola[o]:
            ws.row_dimensions[row_total + o].height = alt_cola[o]
    for o in range(1, 6):
        _pegar(ws, row_total + o, cola[o], {})
        for c, (v, _s) in cola[o].items():           # firmas y textos de pie
            ws.cell(row_total + o, c).value = v
    for m in cola_merges:
        ws.merge_cells(start_row=m.min_row + shift, end_row=m.max_row + shift,
                       start_column=m.min_col, end_column=m.max_col)
    for a, b in [("C", "F"), ("G", "H"), ("I", "L"), ("M", "Q"), ("S", "W")]:
        ws.merge_cells(f"{a}{row_total}:{b}{row_total}")
    ws.print_area = f"B1:AB{row_total + 2}"

    # bloque de cálculo
    r = h0 + shift
    _pegar(ws, r, s_h, {6: "TROPA", 7: "MATERIA PRIMA", 8: "CANTIDAD",
                        9: "KILOGRAMOS", 10: "KG/MEDIA"})
    ws.cell(r, 12, "KG SISTEMA (control)")._style = copy.copy(s_h[9][1])
    r += 1
    for i, t in enumerate(tropas):
        fila_form, ini = FIRST + i, r
        for j, prod in enumerate(PRODUCTOS):
            vals = {7: prod, 8: f"=$M${fila_form}", 9: f"=H{r}*J{r}",
                    10: f"='{HOJA_PESO}'!$D${4 + j}"}
            if j == 0:
                vals[6] = int(t["tropa"]) if t["tropa"].isdigit() else t["tropa"]
            _pegar(ws, r, s_p1 if j == 0 else s_p, vals)
            ws.row_dimensions[r].height = alt_det
            r += 1
        r += 1
        _pegar(ws, r, s_sum, {9: f"=SUM(I{ini}:I{r - 1})"})
        ws.cell(r, 12, d["tot_sis"].get(t["tropa"]))._style = copy.copy(s_sum[9][1])
        ws.row_dimensions[r].height = alt_det
        r += 2

    _hoja_peso(wb, detalle, row_total, ws.title)
    wb.calculation = CalcProperties(fullCalcOnLoad=True)    # Excel recalcula al abrir
    ws.sheet_view.topLeftCell = "A1"                        # abre arriba, en el formulario
    for sel in ws.sheet_view.selection:
        sel.activeCell, sel.sqref = "A1", "A1"
    buf = io.BytesIO()
    wb.save(buf)
    datos = _quitar_custom_height(buf.getvalue(), filas_auto)
    if hasattr(salida, "write"):
        salida.write(datos)
    else:
        with open(salida, "wb") as f:
            f.write(datos)

    res.update(fecha=fecha, fecha_faena=ff,
               serie=m_serie.group(1) if m_serie else None)
    return res


# ----------------------------------------------------------------------- CLI
if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="Genera el CUARTEO final desde la bajada.")
    ap.add_argument("bajada")
    ap.add_argument("plantilla")
    ap.add_argument("salida")
    ap.add_argument("--fecha-faena", type=dt.date.fromisoformat)
    ap.add_argument("--serie", type=int)
    a = ap.parse_args()
    try:
        r = generar(a.bajada, a.plantilla, a.salida, a.fecha_faena, a.serie)
    except BajadaInvalida as e:
        sys.exit(f"Error: {e}")
    print(f"Final generado: {a.salida}  ({r['n_tropas']} tropas, {r['medias']} medias)")
    for av in r["avisos"]:
        print("AVISO:", av, file=sys.stderr)
