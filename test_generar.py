import io
from pathlib import Path

import openpyxl
import pytest

import generar_cuarteo as g

RAIZ = Path(__file__).resolve().parent
BAJADA = RAIZ / "bajada_0610.xlsx"
PLANTILLA = RAIZ / "CUARTEO_plantilla.xlsx"


@pytest.fixture(scope="module")
def generado():
    salida = io.BytesIO()
    res = g.generar(BAJADA, PLANTILLA, salida, serie=2291)
    salida.seek(0)
    return res, openpyxl.load_workbook(salida)


def test_cantidades_y_totales(generado):
    res, _ = generado
    assert res["n_tropas"] == 20
    assert res["medias"] == 1718
    assert res["kg_ingreso"] == 266024


def test_pesos_promedio_desde_la_bajada(generado):
    res, _ = generado
    kg_media = {c["nombre"]: c["kg_media"] for c in res["categorias"]}
    assert kg_media["Asado CV"] == pytest.approx(12.3337, abs=1e-3)
    assert kg_media["Delantero SF"] == pytest.approx(67.2187, abs=1e-3)
    assert kg_media["Entraña Fina"] == pytest.approx(0.2742, abs=1e-3)
    assert kg_media["Falda"] == pytest.approx(7.1882, abs=1e-3)
    assert kg_media["Pistola"] == pytest.approx(60.8672, abs=1e-3)
    assert res["no_va_kg"] == pytest.approx(2541.8, abs=0.1)


def test_fecha_y_serie(generado):
    res, wb = generado
    ws = wb["PD_AnexoIII"]
    assert str(res["fecha"]) == "2026-10-06"
    assert str(res["fecha_faena"]) == "2026-10-05"       # día hábil anterior
    assert res["serie"] == "2291" and ws["P6"].value.endswith("2291")
    assert ws["O14"].value == "Fecha: Día 6 Mes 10 Año 2026"
    assert ws["M38"].value == 1718


def test_detalle_enlazado_a_hoja_de_pesos(generado):
    _, wb = generado
    ws = wb["PD_AnexoIII"]
    assert ws["G45"].value == "Asado CV" and ws["H45"].value == "=$M$18"
    assert ws["J45"].value == "='Peso promedio'!$D$4"
    assert ws["F45"].value == 15770


def test_sin_aviso_de_productos_nuevos(generado):
    res, _ = generado
    assert not [a for a in res["avisos"] if "nuevos" in a]


def test_cantidad_de_tropas_variable(monkeypatch):
    """El formulario se rearma si la bajada trae otra cantidad de tropas."""
    original = g.leer_bajada

    def con_17(src):
        d = original(src)
        d["tropas"] = d["tropas"][:17]
        return d

    monkeypatch.setattr(g, "leer_bajada", con_17)
    salida = io.BytesIO()
    res = g.generar(BAJADA, PLANTILLA, salida, serie=2291)
    ws = openpyxl.load_workbook(io.BytesIO(salida.getvalue()))["PD_AnexoIII"]
    assert res["n_tropas"] == 17
    assert ws["I35"].value == "TOTAL" and ws["M35"].value == res["medias"]
    assert ws.print_area.endswith("$AB$37")
    assert ws["F41"].value == "TROPA" and ws["F42"].value == 15770
    # el total del formulario debe avisar que no coincide con el de la bajada
    assert any("no coincide" in a for a in res["avisos"])


def test_archivo_que_no_es_bajada(tmp_path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "hola"
    p = tmp_path / "x.xlsx"
    wb.save(p)
    with pytest.raises(g.BajadaInvalida):
        g.leer_bajada(p)


def test_producto_nuevo_genera_aviso(tmp_path):
    wb = openpyxl.load_workbook(BAJADA)
    ws = wb.active
    ws["H61"] = "COSTILLAR NUEVO"            # renombra un producto existente
    p = tmp_path / "b.xlsx"
    wb.save(p)
    res = g.generar(p, PLANTILLA, io.BytesIO())
    assert any("COSTILLAR NUEVO" in a for a in res["avisos"])
